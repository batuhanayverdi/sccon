# -*- coding: utf-8 -*-
"""
05_score_people.py  -  Profili cekilen kisileri OpenAI ile kisi bazinda puanlar ve final Excel raporunu uretir.

pip install openai pandas openpyxl
Girdi : data/people_prescored.csv, data/companies_classified.csv, data/raw/profiles/*.json
        (istege bagli) data/interests_map.json  ->  {"int_01": "Smart City", ...}
Cikti : data/SCC_Univention_Report.xlsx
Cache : data/llm_cache/people_<PROMPT_VERSION>.jsonl
"""

import json
import os
import re
import unicodedata
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed

import pandas as pd
from openai import OpenAI
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scc_common import DATA, HERE, RAW, TEAM_SECTOR, print

# -------------------- AYARLAR --------------------
MODEL = "gpt-5-mini"
BATCH = 10
WORKERS = 4
PROMPT_VERSION = "v1"
CACHE = DATA / "llm_cache" / f"people_{PROMPT_VERSION}.jsonl"
REPORT = DATA / "SCC_Univention_Report.xlsx"

KEEP_FOR = [
    "Education Authority / Schulträger (Operator)",
    "Identity/IAM & Directory Services",
    "Platform/SRE/Kubernetes (Ops)",
    "Vendor/Partner (EdTech/SI)",
    "Public Sector IT (Operator/Dienstleister)",
    "Decision Maker (CIO/CDO/IT-Leitung)",
    "Policy/Sponsor",
    "Association/Community/Media",
    "Other (adjacent, DE/DACH)",
    "Not relevant",
]
PRODUCTS = ["UCS", "UCS@school", "Nubus", "openDesk-Kontext"]

SYSTEM = """Du bist Expert:in für den DACH-öffentlichen Sektor (Bund, Länder, Kommunen, Schulträger),
Kubernetes-Platform-Engineering sowie IAM/Directory-Dienste. Ziel: Entscheide, ob eine Person auf der
Smart Country Convention als Gesprächspartner:in / Stakeholder für **Univention** relevant ist:
- **UCS / UCS@school**: Directory/IdM, SSO, Schulserver, App-Ökosystem
- **Nubus**: Kubernetes-basierte IAM-Plattform (Keycloak/OIDC, Multi-Tenant, cloud-native), IAM in openDesk
Univention steht für Open Source und digitale Souveränität.

Du bekommst pro Person: Position, Organisation, die Einordnung der Organisation (sector, org_relevance,
org_why, is_competitor), Rolle auf der Messe (attendee/staff=Aussteller/speaker), Interessen, Sessions.

priority (0-3):
3 = Muss-Gespräch: Entscheider:in oder direkt verantwortlich für IT-Betrieb/IdM/Plattform/Schul-IT
    bei Betreiber, öffentl. IT-Dienstleister oder klar passendem Partner; oder Speaker mit passendem Thema
2 = gutes Gespräch: passende Organisation und plausible Rolle, oder Multiplikator (Verband, Presse, Policy)
1 = nur Kontext: passende Organisation, aber Rolle fachfremd oder unbekannt
0 = nicht relevant
Fehlt die Position, stütze dich auf Organisation, Interessen und Sessions; dann höchstens 2 und confidence "low".

product_fit: nur Produkte mit konkretem Hinweis, sonst leere Liste.
why: 2-3 knappe Punkte auf Deutsch, nur belegbare Fakten aus den Daten, keine Floskeln.
talking_point: EIN konkreter Gesprächseinstieg auf Deutsch (max. 25 Wörter), passend zu Rolle und Organisation.
Wettbewerber (is_competitor=true): priority nach Nutzen für Marktbeobachtung/Partnerschaft, im why erwähnen."""

SCHEMA = {
    "type": "object", "additionalProperties": False, "required": ["results"],
    "properties": {"results": {"type": "array", "items": {
        "type": "object", "additionalProperties": False,
        "required": ["id", "priority", "keep_for", "product_fit", "why", "talking_point", "confidence"],
        "properties": {
            "id": {"type": "string"},
            "priority": {"type": "integer", "enum": [0, 1, 2, 3]},
            "keep_for": {"type": "string", "enum": KEEP_FOR},
            "product_fit": {"type": "array", "items": {"type": "string", "enum": PRODUCTS}},
            "why": {"type": "array", "items": {"type": "string"}},
            "talking_point": {"type": "string"},
            "confidence": {"type": "string", "enum": ["low", "medium", "high"]},
        }}}},
}


# -------------------- YARDIMCILAR --------------------
def get_client() -> OpenAI:
    key = os.environ.get("OPENAI_API_KEY", "").strip()
    kf = HERE / "openai_key.txt"
    if not key and kf.exists():
        key = kf.read_text(encoding="utf-8").strip()
    if not key:
        sys.exit("OpenAI key yok: openai_key.txt olustur ya da OPENAI_API_KEY ver.")
    return OpenAI(api_key=key)


def labels(items, mapping: dict) -> list[str]:
    out = []
    for x in items or []:
        if isinstance(x, dict):
            x = x.get("name") or x.get("label") or x.get("id") or ""
        out.append(mapping.get(str(x), str(x)))
    return [o for o in out if o]


def session_titles(p: dict) -> list[str]:
    out = []
    for x in (p.get("eventDates") or []) + (p.get("events") or []):
        if isinstance(x, dict):
            t = x.get("name") or x.get("title") or x.get("eventName") or ""
            d = x.get("startdate") or x.get("startDate") or x.get("date") or ""
            if t:
                out.append(f"{t} ({d})" if d else t)
        elif x:
            out.append(str(x))
    return list(dict.fromkeys(out))


def load_profiles(imap: dict) -> pd.DataFrame:
    rows = []
    for f in (RAW / "profiles").glob("*.json"):
        p = json.loads(f.read_text(encoding="utf-8")).get("profile", {})
        org = (p.get("organizations") or [{}])[0]
        rows.append({
            "id": f.stem,
            "LinkedIn": p.get("linkedIn") or "",
            "Interests": ", ".join(labels(p.get("interests"), imap)),
            "Offering": ", ".join(labels(p.get("offering"), imap)),
            "LookingFor": ", ".join(labels(p.get("lookingfor"), imap)),
            "Sessions": " | ".join(session_titles(p)),
            "Country": p.get("countrycode") or org.get("country") or "",
            "City": org.get("city") or "",
            "Languages": ", ".join(p.get("languages") or []),
        })
    return pd.DataFrame(rows)


def load_cache() -> dict:
    out = {}
    if CACHE.exists():
        for line in CACHE.read_text(encoding="utf-8").splitlines():
            if line.strip():
                r = json.loads(line)
                out[r["id"]] = r
    return out


_lock = threading.Lock()


def append_cache(rows):
    with _lock:
        CACHE.parent.mkdir(parents=True, exist_ok=True)
        with CACHE.open("a", encoding="utf-8") as f:
            for r in rows:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")


def score_batch(client, batch: pd.DataFrame) -> list[dict]:
    items = [{
        "id": r["id"], "position": r["position"], "organization": r["organization"],
        "sector": r["sector"], "org_relevance": r["relevance"], "org_why": r["why"],
        "is_competitor": bool(r["is_competitor"] == True), "role_at_event": r["userType"],
        "interests": r["Interests"], "offering": r["Offering"], "looking_for": r["LookingFor"],
        "sessions": r["Sessions"],
    } for _, r in batch.iterrows()]
    kwargs = dict(
        model=MODEL,
        messages=[{"role": "system", "content": SYSTEM},
                  {"role": "user", "content": json.dumps(items, ensure_ascii=False, default=str)}],
        response_format={"type": "json_schema",
                         "json_schema": {"name": "people", "strict": True, "schema": SCHEMA}},
    )
    if MODEL.startswith("gpt-5"):
        kwargs["reasoning_effort"] = "low"
    else:
        kwargs["temperature"] = 0
    resp = client.chat.completions.create(**kwargs)
    res = [r for r in json.loads(resp.choices[0].message.content)["results"] if r["id"] in set(batch["id"])]
    print(f"  batch {len(res)}/{len(batch)} ok | tokens in={resp.usage.prompt_tokens} out={resp.usage.completion_tokens}")
    return res


def write_sheet(xw, df: pd.DataFrame, name: str, widths: dict):
    # Excel'in kabul etmedigi gorunmez kontrol karakterlerini temizle (orn. "Rotenburg (W\x02umme)")
    df = df.apply(lambda col: col.map(lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v))
    df.to_excel(xw, sheet_name=name, index=False)
    ws = xw.sheets[name]
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 16)
    wrap = [i for i, col in enumerate(df.columns, 1) if widths.get(col, 16) >= 40]
    for row in ws.iter_rows(min_row=2):
        for i in wrap:
            row[i - 1].alignment = Alignment(wrap_text=True, vertical="top")


# -------------------- ANA AKIS --------------------
def profile_url(name: str, pid: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", unicodedata.normalize("NFKD", name or "")
                  .encode("ascii", "ignore").decode()).strip("-")
    return f"https://online.smartcountry.berlin/person/person-{slug}--u-{pid}"


def main():
    imap_f = DATA / "interests_map.json"
    imap = json.loads(imap_f.read_text(encoding="utf-8")) if imap_f.exists() else {}
    if not imap:
        print("interests_map.json yok, interests kod olarak kalacak.")

    people = pd.read_csv(DATA / "people_prescored.csv").fillna({"position": "", "organization": ""})
    comp = pd.read_csv(DATA / "companies_classified.csv")
    for c in ("is_competitor", "is_exhibitor"):
        comp[c] = comp[c].fillna(False).astype(bool)
    prof = load_profiles(imap)
    print(f"Kisi: {len(people)} | profil dosyasi: {len(prof)}")

    # HERKES: profil yoksa da kisi raporda kalir
    df = people.merge(prof, on="id", how="left")
    df = df.merge(comp[["org_key", "why"]], on="org_key", how="left")
    df = df.fillna({"why": "", "sector": "", "relevance": 0, "Interests": "", "Offering": "",
                    "LookingFor": "", "Sessions": "", "LinkedIn": "", "City": "", "Country": ""})
    df["role_hit"] = df["role_hit"].fillna(False).astype(bool)

    # kurumu alakasiz, rolu eslesmeyen ve speaker olmayanlari API'ye gondermeden 0 say (maliyet)
    auto0 = (df["relevance"] == 0) & ~df["role_hit"] & (df["userType"] != "speaker")
    team = df["sector"] == TEAM_SECTOR   # Univention-Kolleg:innen: nicht bewerten, nur anzeigen

    cache = load_cache()
    todo = df[~auto0 & ~team & ~df["id"].isin(cache)]
    print(f"Otomatik 0: {int(auto0.sum())} | cache: {int(df['id'].isin(cache).sum())} | API'ye gidecek: {len(todo)} | Model: {MODEL}")
    if len(todo):
        client = get_client()
        batches = [todo.iloc[i:i + BATCH] for i in range(0, len(todo), BATCH)]
        with ThreadPoolExecutor(WORKERS) as ex:
            futs = {ex.submit(score_batch, client, b): i for i, b in enumerate(batches)}
            for n, fut in enumerate(as_completed(futs), 1):
                try:
                    rows = fut.result()
                    append_cache(rows)
                    cache.update({r["id"]: r for r in rows})
                except Exception as e:
                    print(f"  !! batch {futs[fut]} hata: {e} (tekrar calistirinca yeniden denenir)")
                if n % 20 == 0:
                    print(f"  ... {n}/{len(batches)} batch")

    sc = pd.DataFrame(cache.values()) if cache else pd.DataFrame(columns=list(SCHEMA["properties"]["results"]["items"]["properties"]))
    sc["product_fit"] = sc["product_fit"].map(lambda x: ", ".join(x) if isinstance(x, list) else "")
    sc["why_person"] = sc["why"].map(lambda w: " • ".join(w) if isinstance(w, list) else "")
    df = df.merge(sc[["id", "priority", "keep_for", "product_fit", "why_person", "talking_point", "confidence"]]
                  .rename(columns={"keep_for": "keep_for_person"}), on="id", how="left")
    df.loc[auto0, ["priority", "keep_for_person", "why_person", "confidence"]] = \
        [0, "Not relevant", "Organisation ohne erkennbaren Bezug (automatisch)", "auto"]
    df.loc[team, ["priority", "keep_for_person", "why_person", "talking_point", "confidence"]] = \
        [0, "Univention-Team", "Kolleg:in von Univention", "", "auto"]

    # LinkedIn verisi (06_linkedin.py ciktisi) varsa ekle
    li_f = DATA / "linkedin.csv"
    if li_f.exists():
        li = pd.read_csv(li_f)
        df = df.merge(li.drop(columns=["LinkedIn"], errors="ignore"), on="id", how="left")
    for c in ("li_followers", "li_headline", "li_about"):
        if c not in df.columns:
            df[c] = ""

    df["ProfileURL"] = [profile_url(n, i) for n, i in zip(df["name"], df["id"])]
    df = df.sort_values(["priority", "relevance", "pre_score"], ascending=False)
    df = df.drop_duplicates(subset=["name", "organization"])  # ayni kisi iki hesapla kayitliysa

    cols = ["priority", "name", "position", "organization", "sector", "userType", "keep_for_person",
            "product_fit", "talking_point", "why_person", "confidence", "Sessions", "Interests",
            "LinkedIn", "li_followers", "li_headline", "is_competitor", "relevance", "City", "Country",
            "ProfileURL"]

    # "NA", "None" gibi kurum adlarini pandas bos (NaN) okuyabiliyor -> metne cevir
    comp["organization"] = comp["organization"].fillna(comp["org_key"]).fillna("").astype(str)
    comp["canonical_org"] = comp["canonical_org"].fillna(comp["organization"]).astype(str)
    org = (comp.groupby("canonical_org")
           .agg(n_people=("n_people", "sum"), exhibitor=("is_exhibitor", "max"),
                speakers=("n_speaker", "sum"), relevance=("relevance", "max"),
                sector=("sector", "first"), is_competitor=("is_competitor", "max"),
                why=("why", "first"), spellings=("organization", lambda x: " | ".join(map(str, x))))
           .reset_index().sort_values(["relevance", "n_people"], ascending=False))
    hot = df[df["priority"] >= 2].groupby("canonical_org").size().rename("n_priority_people")
    org = org.merge(hot, left_on="canonical_org", right_index=True, how="left").fillna({"n_priority_people": 0})

    W = {"name": 22, "position": 32, "organization": 30, "talking_point": 50, "why_person": 60,
         "Sessions": 45, "Interests": 30, "LinkedIn": 35, "keep_for_person": 28, "why": 60,
         "canonical_org": 30, "spellings": 40, "sector": 24, "li_headline": 40, "ProfileURL": 30}
    with pd.ExcelWriter(REPORT, engine="openpyxl") as xw:
        write_sheet(xw, df[df["priority"] >= 2][cols], "meeting_shortlist", W)
        write_sheet(xw, df[cols], "all_attendees", W)
        write_sheet(xw, df[df["userType"] == "speaker"][cols], "speakers", W)
        write_sheet(xw, org[org["relevance"] >= 2], "organizations", W)
        write_sheet(xw, org[org["is_competitor"] == True], "competitors", W)

    print("\npriority dagilimi:")
    print(df["priority"].value_counts(dropna=False).sort_index(ascending=False).to_string())
    print(f"\n-> {REPORT}")


if __name__ == "__main__":
    main()
